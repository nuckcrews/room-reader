import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd

# Load the workbook and select the worksheet
workbook = openpyxl.load_workbook('./example/SaleData.xlsx', data_only=True)
worksheet = workbook.active

# Read the data into a pandas DataFrame
data = pd.DataFrame(worksheet.values)
data.columns = data.iloc[0]
data = data[1:]

# Create a pivot table of the sales data that shows performance by salesman
pivot_table = pd.pivot_table(data, index='SalesMan', values='Sale_amt', aggfunc='sum')

# Create a new sheet in the existing workbook and write the pivot table to it
pivot_sheet = workbook.create_sheet("SalesmanPerformance")

# Write the pivot table to the new sheet
for r in dataframe_to_rows(pivot_table, index=True, header=True):
    pivot_sheet.append(r)

# Create a table style and apply it to the pivot table
table_style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
table = Table(displayName="PivotTable", ref=pivot_sheet.calculate_dimension())
table.tableStyleInfo = table_style
pivot_sheet.add_table(table)

# Save the updated workbook
workbook.save('./example/SaleData.xlsx')
